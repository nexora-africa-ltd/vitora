#!/usr/bin/env python3
"""Generate a Location Decision Matrix Excel workbook."""

import openpyxl
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
    numbers,
)
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule, DataBarRule

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Location Matrix"

# ── Styles ──────────────────────────────────────────────────────────
header_font = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
title_font = Font(name="Calibri", bold=True, size=16, color="1F4E79")
subtitle_font = Font(name="Calibri", bold=True, size=11, color="1F4E79")
factor_font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
weight_font = Font(name="Calibri", bold=True, size=11, color="FFD700")
normal_font = Font(name="Calibri", size=11)
score_font = Font(name="Calibri", bold=True, size=13, color="1F4E79")
rank_font = Font(name="Calibri", bold=True, size=14, color="FFFFFF")

dark_blue_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
medium_blue_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
light_blue_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
light_green_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
light_yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
gold_fill = PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid")
silver_fill = PatternFill(start_color="C0C0C0", end_color="C0C0C0", fill_type="solid")
bronze_fill = PatternFill(start_color="CD7F32", end_color="CD7F32", fill_type="solid")
white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
red_fill = PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid")

thin_border = Border(
    left=Side(style="thin", color="B4C6E7"),
    right=Side(style="thin", color="B4C6E7"),
    top=Side(style="thin", color="B4C6E7"),
    bottom=Side(style="thin", color="B4C6E7"),
)

center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)

# ── Layout constants ────────────────────────────────────────────────
FACTORS = [
    ("Demand (Population + Income)", 0.30),
    ("Accessibility", 0.20),
    ("Competition Gap", 0.20),
    ("Cost (Rent, Setup)", 0.15),
    ("Infrastructure", 0.10),
    ("Regulation", 0.05),
]
NUM_LOCATIONS = 8  # columns for up to 8 locations
FACTOR_START_ROW = 8
LOC_START_COL = 3  # column C onward

# ── Column widths ───────────────────────────────────────────────────
ws.column_dimensions["A"].width = 34
ws.column_dimensions["B"].width = 10
for i in range(NUM_LOCATIONS):
    ws.column_dimensions[get_column_letter(LOC_START_COL + i)].width = 18

# ── Title ───────────────────────────────────────────────────────────
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=LOC_START_COL + NUM_LOCATIONS - 1)
title_cell = ws.cell(row=1, column=1, value="📍 Location Decision Matrix")
title_cell.font = title_font
title_cell.alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[1].height = 36

# ── Instructions ────────────────────────────────────────────────────
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=LOC_START_COL + NUM_LOCATIONS - 1)
ws.cell(row=2, column=1, value="Rate each location 1–10 for every factor. Weighted scores and rankings are calculated automatically.").font = Font(name="Calibri", size=10, italic=True, color="666666")

ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=LOC_START_COL + NUM_LOCATIONS - 1)
ws.cell(row=3, column=1, value="Scale: 1 = Very Poor  |  5 = Average  |  10 = Excellent").font = Font(name="Calibri", size=10, italic=True, color="666666")

# ── Legend row ──────────────────────────────────────────────────────
ROW_LEGEND = 5
ws.merge_cells(start_row=ROW_LEGEND, start_column=1, end_row=ROW_LEGEND, end_column=2)
ws.cell(row=ROW_LEGEND, column=1, value="Score Guide:").font = Font(name="Calibri", bold=True, size=10, color="333333")

legend_items = [("1-3 Poor", "FF6B6B"), ("4-6 Average", "FFF2CC"), ("7-10 Good", "E2EFDA")]
for idx, (label, color) in enumerate(legend_items):
    col = LOC_START_COL + idx
    c = ws.cell(row=ROW_LEGEND, column=col, value=label)
    c.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    c.font = Font(name="Calibri", size=10, bold=True)
    c.alignment = center
    c.border = thin_border

# ── Header row (row 7) ─────────────────────────────────────────────
HEADER_ROW = 7
ws.row_dimensions[HEADER_ROW].height = 32

h1 = ws.cell(row=HEADER_ROW, column=1, value="Factor")
h1.font = header_font; h1.fill = dark_blue_fill; h1.alignment = center; h1.border = thin_border

h2 = ws.cell(row=HEADER_ROW, column=2, value="Weight")
h2.font = header_font; h2.fill = dark_blue_fill; h2.alignment = center; h2.border = thin_border

for i in range(NUM_LOCATIONS):
    col = LOC_START_COL + i
    c = ws.cell(row=HEADER_ROW, column=col, value=f"Location {i + 1}")
    c.font = header_font
    c.fill = medium_blue_fill
    c.alignment = center
    c.border = thin_border

# ── Factor rows ─────────────────────────────────────────────────────
for idx, (factor_name, weight) in enumerate(FACTORS):
    row = FACTOR_START_ROW + idx
    ws.row_dimensions[row].height = 28

    # Factor name
    fc = ws.cell(row=row, column=1, value=factor_name)
    fc.font = Font(name="Calibri", bold=True, size=11, color="1F4E79")
    fc.fill = light_blue_fill
    fc.alignment = left_wrap
    fc.border = thin_border

    # Weight
    wc = ws.cell(row=row, column=2, value=weight)
    wc.font = Font(name="Calibri", bold=True, size=11, color="1F4E79")
    wc.fill = light_blue_fill
    wc.alignment = center
    wc.number_format = "0%"
    wc.border = thin_border

    # Score input cells for each location
    for i in range(NUM_LOCATIONS):
        col = LOC_START_COL + i
        sc = ws.cell(row=row, column=col)
        sc.font = normal_font
        sc.fill = white_fill
        sc.alignment = center
        sc.border = thin_border
        # Data validation 1-10
        from openpyxl.worksheet.datavalidation import DataValidation
        # (added per-cell below)

# ── Data validation (1–10) ──────────────────────────────────────────
dv = DataValidation(
    type="whole", operator="between", formula1=1, formula2=10,
    allow_blank=True,
    showErrorMessage=True,
    errorTitle="Invalid Score",
    error="Please enter a number between 1 and 10.",
    showInputMessage=True,
    promptTitle="Score",
    prompt="Rate this factor 1 (very poor) to 10 (excellent).",
)
for idx in range(len(FACTORS)):
    row = FACTOR_START_ROW + idx
    for i in range(NUM_LOCATIONS):
        col_letter = get_column_letter(LOC_START_COL + i)
        dv.add(f"{col_letter}{row}")
ws.add_data_validation(dv)

# ── Conditional formatting for score cells ──────────────────────────
for idx in range(len(FACTORS)):
    row = FACTOR_START_ROW + idx
    for i in range(NUM_LOCATIONS):
        col_letter = get_column_letter(LOC_START_COL + i)
        cell_ref = f"{col_letter}{row}"
        # Red for 1-3
        ws.conditional_formatting.add(
            cell_ref,
            CellIsRule(operator="between", formula=["1", "3"],
                       fill=PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")),
        )
        # Yellow for 4-6
        ws.conditional_formatting.add(
            cell_ref,
            CellIsRule(operator="between", formula=["4", "6"],
                       fill=light_yellow_fill),
        )
        # Green for 7-10
        ws.conditional_formatting.add(
            cell_ref,
            CellIsRule(operator="between", formula=["7", "10"],
                       fill=light_green_fill),
        )

# ── Separator row ──────────────────────────────────────────────────
SEP_ROW = FACTOR_START_ROW + len(FACTORS)
ws.row_dimensions[SEP_ROW].height = 6
for col in range(1, LOC_START_COL + NUM_LOCATIONS):
    c = ws.cell(row=SEP_ROW, column=col)
    c.fill = dark_blue_fill
    c.border = thin_border

# ── Weighted Score row ──────────────────────────────────────────────
WSCORE_ROW = SEP_ROW + 1
ws.row_dimensions[WSCORE_ROW].height = 32

wsl = ws.cell(row=WSCORE_ROW, column=1, value="WEIGHTED SCORE")
wsl.font = Font(name="Calibri", bold=True, size=12, color="1F4E79")
wsl.fill = light_green_fill
wsl.alignment = center
wsl.border = thin_border

ws.cell(row=WSCORE_ROW, column=2, value="100%").font = Font(name="Calibri", bold=True, size=11, color="1F4E79")
ws.cell(row=WSCORE_ROW, column=2).fill = light_green_fill
ws.cell(row=WSCORE_ROW, column=2).alignment = center
ws.cell(row=WSCORE_ROW, column=2).border = thin_border

for i in range(NUM_LOCATIONS):
    col = LOC_START_COL + i
    col_letter = get_column_letter(col)
    # SUMPRODUCT of (weight * score) for each factor
    parts = []
    for fidx in range(len(FACTORS)):
        frow = FACTOR_START_ROW + fidx
        parts.append(f"$B${frow}*{col_letter}{frow}")
    formula = "=" + "+".join(parts)
    c = ws.cell(row=WSCORE_ROW, column=col, value=formula)
    c.font = score_font
    c.fill = light_green_fill
    c.alignment = center
    c.number_format = "0.00"
    c.border = thin_border

# ── Percentage row (score out of 10 = max) ─────────────────────────
PCT_ROW = WSCORE_ROW + 1
ws.row_dimensions[PCT_ROW].height = 28

ws.cell(row=PCT_ROW, column=1, value="SCORE %").font = Font(name="Calibri", bold=True, size=11, color="1F4E79")
ws.cell(row=PCT_ROW, column=1).fill = light_blue_fill
ws.cell(row=PCT_ROW, column=1).alignment = center
ws.cell(row=PCT_ROW, column=1).border = thin_border

ws.cell(row=PCT_ROW, column=2).fill = light_blue_fill
ws.cell(row=PCT_ROW, column=2).border = thin_border

for i in range(NUM_LOCATIONS):
    col = LOC_START_COL + i
    col_letter = get_column_letter(col)
    # Max possible = 10 (since weights sum to 1.0, max weighted = 10)
    formula = f"=IF({col_letter}{WSCORE_ROW}=0,\"\",{col_letter}{WSCORE_ROW}/10)"
    c = ws.cell(row=PCT_ROW, column=col, value=formula)
    c.font = Font(name="Calibri", bold=True, size=11, color="2E75B6")
    c.fill = light_blue_fill
    c.alignment = center
    c.number_format = "0.0%"
    c.border = thin_border

# ── Rank row ────────────────────────────────────────────────────────
RANK_ROW = PCT_ROW + 1
ws.row_dimensions[RANK_ROW].height = 32

ws.cell(row=RANK_ROW, column=1, value="RANK").font = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
ws.cell(row=RANK_ROW, column=1).fill = dark_blue_fill
ws.cell(row=RANK_ROW, column=1).alignment = center
ws.cell(row=RANK_ROW, column=1).border = thin_border

ws.cell(row=RANK_ROW, column=2).fill = dark_blue_fill
ws.cell(row=RANK_ROW, column=2).border = thin_border

# Build the range of weighted scores for RANK formula
score_cells = ",".join(
    f"{get_column_letter(LOC_START_COL + i)}{WSCORE_ROW}" for i in range(NUM_LOCATIONS)
)

for i in range(NUM_LOCATIONS):
    col = LOC_START_COL + i
    col_letter = get_column_letter(col)
    # RANK descending; show blank if no score
    formula = f'=IF({col_letter}{WSCORE_ROW}=0,"",RANK({col_letter}{WSCORE_ROW},{score_cells}))'
    c = ws.cell(row=RANK_ROW, column=col, value=formula)
    c.font = rank_font
    c.fill = dark_blue_fill
    c.alignment = center
    c.border = thin_border

# ── Recommendation row ──────────────────────────────────────────────
REC_ROW = RANK_ROW + 1
ws.row_dimensions[REC_ROW].height = 28

ws.cell(row=REC_ROW, column=1, value="VERDICT").font = Font(name="Calibri", bold=True, size=11, color="1F4E79")
ws.cell(row=REC_ROW, column=1).fill = light_yellow_fill
ws.cell(row=REC_ROW, column=1).alignment = center
ws.cell(row=REC_ROW, column=1).border = thin_border

ws.cell(row=REC_ROW, column=2).fill = light_yellow_fill
ws.cell(row=REC_ROW, column=2).border = thin_border

for i in range(NUM_LOCATIONS):
    col = LOC_START_COL + i
    col_letter = get_column_letter(col)
    pct_ref = f"{col_letter}{PCT_ROW}"
    formula = (
        f'=IF({pct_ref}="","",IF({pct_ref}>=0.8,"✅ Strong Choice",'
        f'IF({pct_ref}>=0.6,"⚠️ Consider",'
        f'IF({pct_ref}>=0.4,"🔶 Weak","❌ Avoid"))))'
    )
    c = ws.cell(row=REC_ROW, column=col, value=formula)
    c.font = Font(name="Calibri", bold=True, size=10)
    c.fill = light_yellow_fill
    c.alignment = center
    c.border = thin_border

# ── Factor Breakdown sheet ──────────────────────────────────────────
ws2 = wb.create_sheet("Weight Breakdown")
ws2.column_dimensions["A"].width = 34
ws2.column_dimensions["B"].width = 12
ws2.column_dimensions["C"].width = 50

ws2.cell(row=1, column=1, value="Factor").font = header_font
ws2.cell(row=1, column=1).fill = dark_blue_fill
ws2.cell(row=1, column=1).alignment = center
ws2.cell(row=1, column=2, value="Weight").font = header_font
ws2.cell(row=1, column=2).fill = dark_blue_fill
ws2.cell(row=1, column=2).alignment = center
ws2.cell(row=1, column=3, value="What to Assess").font = header_font
ws2.cell(row=1, column=3).fill = dark_blue_fill
ws2.cell(row=1, column=3).alignment = center

guides = [
    ("Demand (Population + Income)", "30%", "Population density, average household income, foot traffic, nearby businesses, target customer presence"),
    ("Accessibility", "20%", "Road access, public transport, parking availability, visibility from main road, ease of finding"),
    ("Competition Gap", "20%", "Number of competitors nearby, unserved niches, market saturation, differentiation opportunity"),
    ("Cost (Rent, Setup)", "15%", "Monthly rent, security deposit, fit-out cost, utility costs, maintenance fees"),
    ("Infrastructure", "10%", "Power reliability, water supply, internet connectivity, waste management, building condition"),
    ("Regulation", "5%", "Zoning compliance, licensing requirements, county permits, health/safety regulations, ease of approval"),
]

for idx, (factor, weight, guide) in enumerate(guides):
    row = idx + 2
    ws2.row_dimensions[row].height = 36
    fill = light_blue_fill if idx % 2 == 0 else white_fill
    for col, val in [(1, factor), (2, weight), (3, guide)]:
        c = ws2.cell(row=row, column=col, value=val)
        c.font = normal_font
        c.fill = fill
        c.alignment = left_wrap
        c.border = thin_border

# ── Print setup ─────────────────────────────────────────────────────
ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
ws.page_setup.fitToWidth = 1
ws.page_setup.orientation = "landscape"

# Freeze panes: freeze factor column + weight column
ws.freeze_panes = "C8"

# ── Save ────────────────────────────────────────────────────────────
output = "/home/thande/dev/vitora/Location_Decision_Matrix.xlsx"
wb.save(output)
print(f"✅ Created: {output}")
