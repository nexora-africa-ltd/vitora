#!/usr/bin/env python3
"""
Generic HTML slide deck to PPTX converter for Vitora decks.

Usage:
    python3 html_to_pptx.py <input.html> [output.pptx]

Parses <section class="slide"> elements and converts them into
PowerPoint slides with proper formatting, tables, and layout.
"""

import sys
import re
from pathlib import Path
from bs4 import BeautifulSoup, NavigableString
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ─── Brand Colors ───
BURGUNDY = RGBColor(0x7A, 0x1B, 0x3E)
TEAL = RGBColor(0x0D, 0x94, 0x88)
TEAL_LIGHT = RGBColor(0x14, 0xB8, 0xA6)
GOLD = RGBColor(0xD4, 0xA8, 0x43)
NAVY = RGBColor(0x0F, 0x17, 0x2A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_TEXT = RGBColor(0xE2, 0xE8, 0xF0)
MUTED_TEXT = RGBColor(0x94, 0xA3, 0xB8)
GREEN = RGBColor(0x22, 0xC5, 0x5E)
RED = RGBColor(0xEF, 0x44, 0x44)
PINK = RGBColor(0xE5, 0x7B, 0xA0)


def get_bg_color(slide_el):
    """Determine background color from slide CSS classes."""
    classes = slide_el.get("class", [])
    if "slide-title" in classes:
        return NAVY
    if "slide-accent" in classes:
        return TEAL
    if "slide-burgundy" in classes:
        return BURGUNDY
    return NAVY


def get_text(el):
    """Get cleaned text from an element, stripping excess whitespace."""
    if el is None:
        return ""
    text = el.get_text(separator=" ", strip=True)
    # Collapse internal whitespace
    text = re.sub(r"\s+", " ", text)
    return text


def get_multiline_text(el):
    """Get text preserving <br> as newlines."""
    if el is None:
        return ""
    # Replace <br> with newlines
    for br in el.find_all("br"):
        br.replace_with("\n")
    text = el.get_text(strip=False)
    # Clean up but preserve intentional newlines
    lines = [line.strip() for line in text.split("\n")]
    return "\n".join(line for line in lines if line)


def set_slide_bg(slide, color):
    """Set solid background color."""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text_box(slide, left, top, width, height, text, font_size=18,
                 bold=False, color=WHITE, alignment=PP_ALIGN.LEFT):
    """Add a text box."""
    if not text:
        return None
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = "Calibri"
    p.alignment = alignment
    return txBox


def add_paragraph(tf, text, font_size=14, bold=False, color=LIGHT_TEXT,
                  alignment=PP_ALIGN.LEFT, space_before=Pt(6)):
    """Add a paragraph to an existing text frame."""
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = "Calibri"
    p.alignment = alignment
    if space_before:
        p.space_before = space_before
    return p


def add_table_to_slide(slide, table_el, left, top, width):
    """Convert an HTML table to a PPTX table."""
    rows_data = []
    header_data = []

    thead = table_el.find("thead")
    if thead:
        for tr in thead.find_all("tr"):
            cells = [get_text(th) for th in tr.find_all(["th", "td"])]
            if cells:
                header_data.append(cells)

    tbody = table_el.find("tbody")
    if tbody:
        for tr in tbody.find_all("tr"):
            # Skip empty spacer rows
            cells = tr.find_all(["td", "th"])
            if not cells:
                continue
            # Skip colspan spacer rows
            if len(cells) == 1 and cells[0].get("colspan"):
                text = get_text(cells[0])
                if not text or text.strip() == "":
                    continue
                # Add as a note row
                rows_data.append([text])
                continue
            row = [get_text(cell) for cell in cells]
            if any(row):
                rows_data.append(row)
    else:
        for tr in table_el.find_all("tr"):
            cells = tr.find_all(["td", "th"])
            row = [get_text(cell) for cell in cells]
            if any(row):
                rows_data.append(row)

    all_rows = header_data + rows_data
    if not all_rows:
        return top

    # Determine column count (use max across all rows)
    num_cols = max(len(r) for r in all_rows)
    num_rows = len(all_rows)

    # Calculate dimensions
    row_height = Inches(0.38)
    total_height = row_height * num_rows
    col_width = width // num_cols

    # Cap table height
    max_height = Inches(4.5)
    if total_height > max_height:
        row_height = max_height // num_rows
        total_height = max_height

    tbl_shape = slide.shapes.add_table(num_rows, num_cols, left, top, width, total_height)
    tbl = tbl_shape.table

    # Style table
    for i, row in enumerate(all_rows):
        for j in range(num_cols):
            cell = tbl.cell(i, j)
            cell_text = row[j] if j < len(row) else ""
            cell.text = cell_text

            # Format
            for para in cell.text_frame.paragraphs:
                para.font.size = Pt(9)
                para.font.name = "Calibri"
                if i < len(header_data):
                    para.font.bold = True
                    para.font.color.rgb = MUTED_TEXT
                else:
                    para.font.color.rgb = LIGHT_TEXT
                # Right-align numeric columns (not first column)
                if j > 0:
                    para.alignment = PP_ALIGN.RIGHT
                else:
                    para.alignment = PP_ALIGN.LEFT

            # Cell fill
            cell.fill.solid()
            if i < len(header_data):
                cell.fill.fore_color.rgb = RGBColor(0x1A, 0x24, 0x35)
            else:
                cell.fill.fore_color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    return top + total_height + Inches(0.2)


def process_cards(slide, cards, start_y, max_cols=3):
    """Render card elements as rounded rectangles."""
    num_cards = len(cards)
    cols = min(num_cards, max_cols)
    card_width = Inches(12.0 / cols)
    gap = Inches(0.15)

    for i, card in enumerate(cards):
        col = i % cols
        row = i // cols
        x = Inches(0.5) + col * (card_width + gap)
        y = start_y + row * Inches(2.0)

        h4 = card.find("h4")
        title = get_text(h4) if h4 else ""
        # Get paragraphs/descriptions
        paras = card.find_all("p")
        desc = get_text(paras[0]) if paras else ""

        # Get assumption items
        items = card.find_all("div", class_="assumption-item")
        item_lines = []
        for item in items:
            label_el = item.find("span", class_="label")
            value_el = item.find("span", class_="value")
            if label_el and value_el:
                item_lines.append(f"{get_text(label_el)}: {get_text(value_el)}")

        # Fallback: if no h4/p/assumption-items found, extract all direct child
        # div text content (handles cards with plain <div> content like stat cards)
        extra_lines = []
        if not title and not desc and not item_lines:
            for child in card.find_all("div", recursive=False):
                # Skip assumption-list containers
                if child.get("class") and "assumption-list" in child.get("class", []):
                    continue
                child_text = get_text(child)
                if child_text:
                    extra_lines.append(child_text)

        # Determine card height
        content_lines = (1 if title else 0) + (1 if desc else 0) + len(item_lines) + len(extra_lines)
        card_height = Inches(max(1.4, 0.4 + content_lines * 0.3))

        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_width - gap, card_height
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        shape.line.color.rgb = RGBColor(0x33, 0x41, 0x55)

        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Pt(12)
        tf.margin_top = Pt(10)

        if title:
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(12)
            p.font.bold = True
            p.font.color.rgb = LIGHT_TEXT

        if desc:
            add_paragraph(tf, desc, font_size=10, color=MUTED_TEXT, space_before=Pt(4))

        for line in item_lines:
            add_paragraph(tf, line, font_size=9, color=LIGHT_TEXT, space_before=Pt(3))

        # Render fallback div text lines (for cards without h4/p structure)
        for idx, line in enumerate(extra_lines):
            # First line is usually the label, second is the big number, third is subtitle
            if idx == 0:
                add_paragraph(tf, line, font_size=9, bold=True, color=MUTED_TEXT,
                              alignment=PP_ALIGN.CENTER, space_before=Pt(6))
            elif idx == 1:
                add_paragraph(tf, line, font_size=22, bold=True, color=TEAL_LIGHT,
                              alignment=PP_ALIGN.CENTER, space_before=Pt(4))
            else:
                add_paragraph(tf, line, font_size=9, color=MUTED_TEXT,
                              alignment=PP_ALIGN.CENTER, space_before=Pt(2))

    total_rows = (num_cards + cols - 1) // cols
    return start_y + total_rows * Inches(2.0) + Inches(0.2)


def process_kpi_chips(slide, chips, y_pos):
    """Render KPI chips as a row."""
    num = len(chips)
    if num == 0:
        return y_pos
    chip_width = Inches(12.0 / num)

    for i, chip in enumerate(chips):
        val_el = chip.find(class_="val")
        lbl_el = chip.find(class_="lbl")
        val = get_text(val_el) if val_el else ""
        lbl = get_text(lbl_el) if lbl_el else ""

        x = Inches(0.5) + i * chip_width
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, x, y_pos, chip_width - Inches(0.1), Inches(0.8)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        shape.line.color.rgb = RGBColor(0x33, 0x41, 0x55)

        tf = shape.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = val
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = TEAL_LIGHT
        p.alignment = PP_ALIGN.CENTER
        add_paragraph(tf, lbl, font_size=9, color=MUTED_TEXT,
                      alignment=PP_ALIGN.CENTER, space_before=Pt(2))

    return y_pos + Inches(1.0)


def process_stat_blocks(slide, blocks, y_pos):
    """Render stat-block elements."""
    num = len(blocks)
    if num == 0:
        return y_pos
    cols = min(num, 4)
    block_width = Inches(12.0 / cols)

    for i, block in enumerate(blocks):
        col = i % cols
        row = i // cols
        x = Inches(0.5) + col * block_width
        y = y_pos + row * Inches(1.6)

        num_el = block.find(class_="stat-number")
        lbl_el = block.find(class_="stat-label")
        val = get_text(num_el) if num_el else ""
        lbl = get_text(lbl_el) if lbl_el else ""

        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, x, y, block_width - Inches(0.15), Inches(1.4)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        shape.line.color.rgb = RGBColor(0x33, 0x41, 0x55)

        tf = shape.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = val
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = TEAL_LIGHT
        p.alignment = PP_ALIGN.CENTER
        add_paragraph(tf, lbl, font_size=9, bold=True, color=MUTED_TEXT,
                      alignment=PP_ALIGN.CENTER, space_before=Pt(4))

    total_rows = (num + cols - 1) // cols
    return y_pos + total_rows * Inches(1.6) + Inches(0.2)


def process_bar_chart(slide, chart_el, y_pos):
    """Render bar chart rows."""
    rows = chart_el.find_all("div", class_="bar-row")
    for i, row in enumerate(rows):
        label_el = row.find(class_="bar-label")
        value_el = row.find(class_="bar-value")
        fill_el = row.find(class_="bar-fill")

        label = get_text(label_el) if label_el else ""
        value = get_text(value_el) if value_el else ""
        fill_text = get_text(fill_el) if fill_el else ""

        y = y_pos + i * Inches(0.55)
        add_text_box(slide, Inches(0.6), y, Inches(1.2), Inches(0.4),
                     label, font_size=11, bold=True, color=LIGHT_TEXT)
        add_text_box(slide, Inches(1.9), y, Inches(5.5), Inches(0.4),
                     fill_text, font_size=10, color=LIGHT_TEXT)
        add_text_box(slide, Inches(7.5), y, Inches(2.0), Inches(0.4),
                     value, font_size=12, bold=True, color=WHITE, alignment=PP_ALIGN.RIGHT)

    return y_pos + len(rows) * Inches(0.55) + Inches(0.2)


def convert_slide(prs, slide_el, slide_num):
    """Convert a single HTML slide section to a PPTX slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    bg_color = get_bg_color(slide_el)
    set_slide_bg(slide, bg_color)

    # Track vertical position
    y_cursor = Inches(0.8)

    # Section label
    label_el = slide_el.find(class_="section-label")
    if label_el:
        label_color = TEAL_LIGHT
        if "slide-accent" in slide_el.get("class", []):
            label_color = RGBColor(0xCC, 0xFC, 0xE8)
        elif "slide-burgundy" in slide_el.get("class", []):
            label_color = RGBColor(0xFC, 0xD3, 0xDB)
        add_text_box(slide, Inches(0.6), y_cursor, Inches(4), Inches(0.3),
                     get_text(label_el).upper(), font_size=10, bold=True, color=label_color)
        y_cursor += Inches(0.35)

    # Title (h1 or h2)
    title_el = slide_el.find("h1") or slide_el.find("h2")
    if title_el:
        title_text = get_multiline_text(title_el)
        is_title_slide = "slide-title" in slide_el.get("class", [])
        font_size = 48 if is_title_slide else 32
        alignment = PP_ALIGN.CENTER if is_title_slide else PP_ALIGN.LEFT
        top = Inches(2.2) if is_title_slide else y_cursor

        add_text_box(slide, Inches(0.6), top, Inches(12.0), Inches(1.2),
                     title_text, font_size=font_size, bold=True, color=WHITE,
                     alignment=alignment)
        y_cursor = top + Inches(1.3)

    # Title badge (for title slides)
    badge_el = slide_el.find(class_="title-badge")
    if badge_el:
        add_text_box(slide, Inches(3.5), Inches(1.6), Inches(6), Inches(0.4),
                     get_text(badge_el), font_size=11, bold=True, color=TEAL_LIGHT,
                     alignment=PP_ALIGN.CENTER)

    # Subtitle
    subtitle_el = slide_el.find(class_="subtitle") or slide_el.find(class_="title-sub")
    if subtitle_el:
        sub_color = MUTED_TEXT
        if "slide-accent" in slide_el.get("class", []):
            sub_color = RGBColor(0xCC, 0xFC, 0xE8)
        elif "slide-burgundy" in slide_el.get("class", []):
            sub_color = RGBColor(0xFC, 0xD3, 0xDB)
        elif "slide-title" in slide_el.get("class", []):
            sub_color = LIGHT_TEXT

        is_title_slide = "slide-title" in slide_el.get("class", [])
        alignment = PP_ALIGN.CENTER if is_title_slide else PP_ALIGN.LEFT

        add_text_box(slide, Inches(0.6), y_cursor, Inches(10), Inches(0.7),
                     get_text(subtitle_el), font_size=13, color=sub_color,
                     alignment=alignment)
        y_cursor += Inches(0.8)

    # Process content elements in order
    # Find tables
    tables = slide_el.find_all("table", class_="data-table")
    for table in tables:
        y_cursor = add_table_to_slide(slide, table, Inches(0.5), y_cursor, Inches(12.0))

    # Find stat grids
    stat_grids = slide_el.find_all(class_="stat-grid")
    for grid in stat_grids:
        blocks = grid.find_all(class_="stat-block")
        y_cursor = process_stat_blocks(slide, blocks, y_cursor)

    # Find standalone stat blocks (not in stat-grid)
    standalone_stats = []
    for block in slide_el.find_all(class_="stat-block"):
        parent = block.parent
        if not parent or "stat-grid" not in parent.get("class", []):
            standalone_stats.append(block)
    if standalone_stats:
        y_cursor = process_stat_blocks(slide, standalone_stats, y_cursor)

    # Find cards (not inside other cards or grids already processed)
    card_containers = slide_el.find_all(class_=re.compile(r"grid-[234]"))
    for container in card_containers:
        cards = container.find_all("div", class_="card", recursive=False)
        if cards:
            cols = 3
            cls = container.get("class", [])
            if "grid-4" in cls:
                cols = 4
            elif "grid-2" in cls:
                cols = 2
            y_cursor = process_cards(slide, cards, y_cursor, max_cols=cols)

    # Standalone cards not in grids (e.g. in two-col layouts)
    two_cols = slide_el.find_all(class_="two-col")
    for two_col in two_cols:
        cards = two_col.find_all("div", class_="card", recursive=False)
        if cards:
            y_cursor = process_cards(slide, cards, y_cursor, max_cols=2)

    # KPI chips
    kpi_rows = slide_el.find_all(class_="kpi-row")
    for kpi_row in kpi_rows:
        chips = kpi_row.find_all(class_="kpi-chip")
        y_cursor = process_kpi_chips(slide, chips, y_cursor)

    # Bar charts
    bar_charts = slide_el.find_all(class_="bar-chart")
    for chart in bar_charts:
        y_cursor = process_bar_chart(slide, chart, y_cursor)

    # Assumption lists (standalone, not in cards)
    for assumption_list in slide_el.find_all(class_="assumption-list"):
        parent = assumption_list.parent
        if parent and "card" in parent.get("class", []):
            continue  # Already handled via cards
        items = assumption_list.find_all(class_="assumption-item")
        for item in items:
            label_el = item.find(class_="label")
            value_el = item.find(class_="value")
            if label_el and value_el:
                text = f"{get_text(label_el)}: {get_text(value_el)}"
                add_text_box(slide, Inches(0.8), y_cursor, Inches(8), Inches(0.35),
                             text, font_size=11, color=LIGHT_TEXT)
                y_cursor += Inches(0.35)

    # Footnote
    footnote_el = slide_el.find(class_="footnote")
    if footnote_el:
        add_text_box(slide, Inches(0.6), Inches(6.9), Inches(11), Inches(0.4),
                     get_text(footnote_el), font_size=9, color=MUTED_TEXT)

    # Logo
    add_text_box(slide, Inches(0.4), Inches(0.2), Inches(2), Inches(0.35),
                 "● VITORA", font_size=12, bold=True, color=MUTED_TEXT)

    # Slide number
    add_text_box(slide, Inches(12.3), Inches(7.0), Inches(0.8), Inches(0.3),
                 f"{slide_num:02d}", font_size=10, color=MUTED_TEXT, alignment=PP_ALIGN.RIGHT)


def convert_html_to_pptx(input_path, output_path=None):
    """Main conversion function."""
    input_path = Path(input_path)
    if output_path is None:
        output_path = input_path.with_suffix(".pptx")
    else:
        output_path = Path(output_path)

    # Read and parse HTML
    html_content = input_path.read_text(encoding="utf-8")
    soup = BeautifulSoup(html_content, "lxml")

    # Find all slide sections
    slides = soup.find_all("section", class_="slide")
    if not slides:
        print(f"❌ No <section class='slide'> elements found in {input_path}")
        sys.exit(1)

    print(f"📄 Found {len(slides)} slides in {input_path.name}")

    # Create presentation
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    # Convert each slide
    for i, slide_el in enumerate(slides, 1):
        print(f"  → Slide {i:02d}...", end=" ")
        convert_slide(prs, slide_el, i)
        print("✓")

    # Save
    prs.save(str(output_path))
    print(f"\n✅ Saved: {output_path}")
    return output_path


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python3 html_to_pptx.py <input.html> [output.pptx]")
        sys.exit(1)

    input_file = sys.argv[1]
    output_file = sys.argv[2] if len(sys.argv) > 2 else None
    convert_html_to_pptx(input_file, output_file)
